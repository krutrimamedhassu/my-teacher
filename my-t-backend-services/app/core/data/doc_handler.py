import os
from pathlib import Path
from typing import Dict, Any, Optional, List, Union
import json
from enum import Enum
from app.logger.app_logger import app_logger as logger
from app.core.config import settings

from PyPDF2 import PdfReader
from docx import Document as DocxDocument

class DocumentType(str, Enum):
    PDF = "pdf"
    TEXT = "txt"
    DOCX = "docx"
    HTML = "html"
    UNKNOWN = "unknown"

class Document:
    def __init__(self, content: str, doc_type: DocumentType,
                 filename: Optional[str] = None, metadata: Optional[Dict[str, Any]] = None):
        logger.log_debug(f"[Document] Initializing document with type: {doc_type}, filename: {filename}")
        self.content = content
        self.doc_type = doc_type
        self.filename = filename
        self.metadata = metadata or {}
        logger.log_debug(f"[Document] Document initialized successfully with content length: {len(content)}")

    def to_dict(self) -> Dict[str, Any]:
        logger.log_debug(f"[Document] Converting document to dict: {self.filename}")
        try:
            result = {
                "content": self.content,
                "doc_type": self.doc_type.value,
                "filename": self.filename,
                "metadata": self.metadata
            }
            logger.log_debug(f"[Document] Successfully converted document to dict")
            return result
        except Exception as e:
            logger.log_error(f"[Document] Error in to_dict: {str(e)}")
            raise

    @classmethod
    def from_dict(cls, data: Dict[str, Any]) -> "Document":
        logger.log_debug(f"[Document] Creating document from dict for filename: {data.get('filename')}")
        try:
            document = cls(
                content=data["content"],
                doc_type=DocumentType(data["doc_type"]),
                filename=data.get("filename"),
                metadata=data.get("metadata")
            )
            logger.log_debug(f"[Document] Successfully created document from dict")
            return document
        except Exception as e:
            logger.log_error(f"[Document] Error in from_dict: {str(e)}")
            raise

    def save(self, path: Union[str, Path]) -> None:
        logger.log_info(f"[Document] Starting save to path: {path}")
        try:
            path = Path(path)
            with open(path, "w", encoding="utf-8") as f:
                json.dump(self.to_dict(), f, indent=2)
            logger.log_info(f"[Document] Successfully saved document to {path}")
        except Exception as e:
            logger.log_error(f"[Document] Error saving document: {str(e)}")
            raise

    @classmethod
    def load(cls, path: Union[str, Path]) -> "Document":
        logger.log_info(f"[Document] Starting load from path: {path}")
        try:
            path = Path(path)
            with open(path, "r", encoding="utf-8") as f:
                data = json.load(f)
            document = cls.from_dict(data)
            logger.log_info(f"[Document] Successfully loaded document from {path}")
            return document
        except Exception as e:
            logger.log_error(f"[Document] Error loading document: {str(e)}")
            raise

class DocumentHandler:
    @staticmethod
    def get_document_type(filename: str) -> DocumentType:
        logger.log_debug(f"[DocumentHandler] Getting document type for filename: {filename}")
        try:
            ext = filename.split(".")[-1].lower()
            if ext in [t.value for t in DocumentType]:
                doc_type = DocumentType(ext)
                logger.log_debug(f"[DocumentHandler] Detected document type: {doc_type} for {filename}")
                return doc_type
            logger.log_warning(f"[DocumentHandler] Unknown document type for extension: {ext}")
            return DocumentType.UNKNOWN
        except Exception as e:
            logger.log_error(f"[DocumentHandler] Error determining document type: {str(e)}")
            return DocumentType.UNKNOWN

    @staticmethod
    def validate_document(file_path: Union[str, Path], max_size_mb: Optional[int] = None) -> bool:
        logger.log_info(f"[DocumentHandler] Starting validate_document for: {file_path}")
        try:
            max_size_mb = max_size_mb or settings.MAX_DOCUMENT_SIZE_MB
            path = Path(file_path)
            if not path.exists():
                logger.log_error(f"[DocumentHandler] File not found: {file_path}")
                return False
            file_size_mb = path.stat().st_size / (1024 * 1024)
            if file_size_mb > max_size_mb:
                logger.log_error(f"[DocumentHandler] File too large: {file_size_mb:.2f}MB (max: {max_size_mb}MB)")
                return False
            doc_type = DocumentHandler.get_document_type(path.name)
            if doc_type == DocumentType.UNKNOWN:
                logger.log_error(f"[DocumentHandler] Unsupported file type: {path.suffix}")
                return False
            logger.log_info(f"[DocumentHandler] Document validation successful for {file_path}")
            return True
        except Exception as e:
            logger.log_error(f"[DocumentHandler] Error in validate_document: {str(e)}")
            return False

    @staticmethod
    def parse_pdf(file_path: Union[str, Path]) -> Document:
        logger.log_info(f"[DocumentHandler] Starting parse_pdf for: {file_path}")
        try:
            reader = PdfReader(str(file_path))
            logger.log_debug(f"[DocumentHandler] PDF has {len(reader.pages)} pages")
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
            metadata = {"pages": len(reader.pages)}
            document = Document(
                content=text,
                doc_type=DocumentType.PDF,
                filename=Path(file_path).name,
                metadata=metadata
            )
            logger.log_info(f"[DocumentHandler] Successfully parsed PDF with {len(text)} characters")
            return document
        except Exception as e:
            logger.log_error(f"[DocumentHandler] PDF parsing failed: {str(e)}")
            raise

    @staticmethod
    def parse_text(file_path: Union[str, Path]) -> Document:
        logger.log_info(f"[DocumentHandler] Starting parse_text for: {file_path}")
        try:
            with open(file_path, "r", encoding="utf-8") as f:
                content = f.read()
            document = Document(
                content=content,
                doc_type=DocumentType.TEXT,
                filename=Path(file_path).name
            )
            logger.log_info(f"[DocumentHandler] Successfully parsed text file with {len(content)} characters")
            return document
        except Exception as e:
            logger.log_error(f"[DocumentHandler] Text parsing failed: {str(e)}")
            raise

    @staticmethod
    def parse_docx(file_path: Union[str, Path]) -> Document:
        logger.log_info(f"[DocumentHandler] Starting parse_docx for: {file_path}")
        try:
            doc = DocxDocument(str(file_path))
            logger.log_debug(f"[DocumentHandler] DOCX has {len(doc.paragraphs)} paragraphs")
            content = "\n".join([p.text for p in doc.paragraphs])
            metadata = {"paragraphs": len(doc.paragraphs)}
            document = Document(
                content=content,
                doc_type=DocumentType.DOCX,
                filename=Path(file_path).name,
                metadata=metadata
            )
            logger.log_info(f"[DocumentHandler] Successfully parsed DOCX with {len(content)} characters")
            return document
        except Exception as e:
            logger.log_error(f"[DocumentHandler] DOCX parsing failed: {str(e)}")
            raise

    @staticmethod
    def parse_html(html_content: str, url: Optional[str] = None) -> Document:
        logger.log_info(f"[DocumentHandler] Starting parse_html from {url or 'unknown source'}")
        try:
            from bs4 import BeautifulSoup
            soup = BeautifulSoup(html_content, "html.parser")
            logger.log_debug(f"[DocumentHandler] HTML content parsed with BeautifulSoup")
            # Extract visible text (skip scripts/styles)
            for script in soup(["script", "style"]):
                script.extract()
            content = soup.get_text(separator="\n")
            content = "\n".join(line.strip() for line in content.splitlines() if line.strip())
            document = Document(
                content=content,
                doc_type=DocumentType.HTML,
                filename=url.split("/")[-1] if url else "webpage.html",
                metadata={"url": url}
            )
            logger.log_info(f"[DocumentHandler] Successfully parsed HTML with {len(content)} characters")
            return document
        except Exception as e:
            logger.log_error(f"[DocumentHandler] HTML parsing failed: {str(e)}")
            raise

    @staticmethod
    def parse_document(file_path: Union[str, Path]) -> Document:
        logger.log_info(f"[DocumentHandler] Starting parse_document for: {file_path}")
        try:
            if not DocumentHandler.validate_document(file_path):
                logger.log_error(f"[DocumentHandler] Document validation failed for: {file_path}")
                raise ValueError(f"Invalid document: {file_path}")
            doc_type = DocumentHandler.get_document_type(Path(file_path).name)
            logger.log_debug(f"[DocumentHandler] Determined document type: {doc_type}")
            if doc_type == DocumentType.PDF:
                document = DocumentHandler.parse_pdf(file_path)
            elif doc_type == DocumentType.TEXT:
                document = DocumentHandler.parse_text(file_path)
            elif doc_type == DocumentType.DOCX:
                document = DocumentHandler.parse_docx(file_path)
            else:
                logger.log_error(f"[DocumentHandler] Unsupported document type: {doc_type}")
                raise ValueError(f"Unsupported document type: {doc_type}")
            logger.log_info(f"[DocumentHandler] Successfully parsed document: {file_path}")
            return document
        except Exception as e:
            logger.log_error(f"[DocumentHandler] Error in parse_document: {str(e)}")
            raise

    @staticmethod
    def parse_url(url: str) -> Document:
        logger.log_info(f"[DocumentHandler] Starting parse_url for: {url}")
        try:
            import requests
            logger.log_debug(f"[DocumentHandler] Fetching content from URL: {url}")
            resp = requests.get(url)
            resp.raise_for_status()
            html_content = resp.text
            logger.log_debug(f"[DocumentHandler] Received {len(html_content)} characters from URL")
            document = DocumentHandler.parse_html(html_content, url)
            logger.log_info(f"[DocumentHandler] Successfully parsed document from URL: {url}")
            return document
        except Exception as e:
            logger.log_error(f"[DocumentHandler] Error fetching/parsing URL {url}: {str(e)}")
            raise

    @staticmethod
    def extract_text_chunks(document: Document, chunk_size: int = 1000, overlap: int = 200) -> List[str]:
        logger.log_info(f"[DocumentHandler] Starting extract_text_chunks from document: {document.filename}")
        try:
            content = document.content or ""
            logger.log_debug(f"[DocumentHandler] Document content length: {len(content)} characters")
            if len(content) <= chunk_size:
                logger.log_info(f"[DocumentHandler] Document fits in single chunk, returning as-is")
                return [content]
            chunks = []
            start = 0
            while start < len(content):
                end = min(start + chunk_size, len(content))
                # Try to split at paragraph or sentence boundaries
                if end < len(content):
                    paragraph_break = content.rfind("\n\n", start, end)
                    if paragraph_break != -1 and paragraph_break > start + chunk_size // 2:
                        end = paragraph_break + 2
                    else:
                        sentence_break = max(
                            content.rfind(". ", start, end),
                            content.rfind("! ", start, end),
                            content.rfind("? ", start, end),
                        )
                        if sentence_break != -1 and sentence_break > start + chunk_size // 2:
                            end = sentence_break + 2
                chunks.append(content[start:end].strip())
                start = end - overlap
            logger.log_info(f"[DocumentHandler] Successfully extracted {len(chunks)} chunks from document")
            return chunks
        except Exception as e:
            logger.log_error(f"[DocumentHandler] Error in extract_text_chunks: {str(e)}")
            raise